"""Opt-in XLSX importer for ordinary tabular sheets.

Uses the same destination table as the legacy importer but never updates existing
records. Preview is mandatory in the UI; callers must explicitly confirm writes.
"""
import re
import unicodedata
from datetime import date, datetime
from flask import jsonify, request
from flask_login import login_required, current_user

ALIASES = {
    'nome': {'nome', 'nome completo', 'aluno', 'nome aluno', 'nome do aluno', 'estagiario', 'residente'},
    'especialidade': {'especialidade', 'area', 'area medica', 'servico'},
    'email': {'email', 'e mail', 'correio eletronico'},
    'telefone': {'telefone', 'celular', 'whatsapp', 'fone'},
    'mes_ano': {'mes ano', 'competencia', 'mes de referencia'},
    'termino': {'termino', 'data termino', 'data fim', 'fim'},
    'cracha': {'cracha', 'numero cracha'},
    'observacao': {'observacao', 'observacoes', 'notas'},
    'tipo_id': {'tipo', 'modalidade', 'tipo estagio'},
    'semana': {'semana'},
}
LOOKUP = {alias: field for field, names in ALIASES.items() for alias in names}


def normalized(value):
    text = unicodedata.normalize('NFKD', str(value or ''))
    text = ''.join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r'[^a-z0-9]+', ' ', text.casefold()).strip()


def detect_header(rows):
    for index, row in enumerate(rows[:20]):
        mapping = {LOOKUP[normalized(v)]: i for i, v in enumerate(row)
                   if normalized(v) in LOOKUP}
        if 'nome' in mapping and 'especialidade' in mapping:
            return index, mapping
    return None, {}


def parse_sheet(ws, sheet_name, max_rows=5000):
    iterator = ws.iter_rows(values_only=True)
    rows = []
    for _, row in zip(range(20), iterator):
        rows.append(row)
    offset, mapping = detect_header(rows)
    if offset is None:
        return [], [{'aba': sheet_name, 'linha': 0, 'motivo': 'Cabecalho Nome/Especialidade nao encontrado'}]
    records, errors = [], []
    def convert(row, line):
        fields = {field: row[index] if index < len(row) else None for field, index in mapping.items()}
        nome = str(fields.get('nome') or '').strip()
        specialty = str(fields.get('especialidade') or '').strip()
        if not nome and not specialty:
            return None
        if not nome or not specialty:
            errors.append({'aba': sheet_name, 'linha': line, 'motivo': 'Nome e especialidade obrigatorios'})
            return None
        raw_month = fields.get('mes_ano')
        if isinstance(raw_month, (date, datetime)):
            month = raw_month.strftime('%Y-%m')
        elif raw_month is None or not str(raw_month).strip():
            month = None
        else:
            raw = str(raw_month).strip()
            match = re.fullmatch(r'(20\d{2})[-/](0?[1-9]|1[0-2])', raw)
            if not match:
                errors.append({'aba': sheet_name, 'linha': line, 'motivo': 'Mes/ano invalido: use AAAA-MM'})
                return None
            month = f'{match.group(1)}-{int(match.group(2)):02d}'
        raw_week = fields.get('semana')
        try:
            week = int(raw_week) if raw_week is not None and str(raw_week).strip() else 1
            if not 1 <= week <= 6:
                raise ValueError()
        except (ValueError, TypeError):
            errors.append({'aba': sheet_name, 'linha': line, 'motivo': 'Semana invalida (1 a 6)'})
            return None
        raw_type = normalized(fields.get('tipo_id'))
        kind = 2 if 'obrig' in raw_type else 3 if 'opta' in raw_type else 1
        end = fields.get('termino')
        if isinstance(end, (date, datetime)):
            end = end.strftime('%Y-%m-%d')
        elif end:
            end = str(end).strip()
            try:
                date.fromisoformat(end)
            except ValueError:
                errors.append({'aba': sheet_name, 'linha': line, 'motivo': 'Termino invalido: use AAAA-MM-DD'})
                return None
        return {
            'nome': nome, 'especialidade': specialty, 'mes_ano': month,
            'semana': week, 'tipo_id': kind, 'etapa': 1 if kind == 1 else 0,
            'termino': end or None, 'cracha': str(fields['cracha']).strip() if fields.get('cracha') is not None else None,
            'email': str(fields['email']).strip() if fields.get('email') else None,
            'telefone': str(fields['telefone']).strip() if fields.get('telefone') else None,
            'observacao': str(fields['observacao']).strip() if fields.get('observacao') else None,
        }
    for idx, row in enumerate(rows[offset + 1:], offset + 2):
        item = convert(row, idx)
        if item:
            records.append(item)
    for idx, row in enumerate(iterator, len(rows) + 1):
        if idx > max_rows:
            errors.append({'aba': sheet_name, 'linha': idx, 'motivo': 'Limite de 5000 linhas por aba excedido'})
            break
        item = convert(row, idx)
        if item:
            records.append(item)
    return records, errors


def key_of(record):
    return (normalized(record['nome']), normalized(record['especialidade']), record.get('mes_ano') or '')


def register_flexible_import(app, get_db):
    @app.route('/api/importar-flexivel', methods=['POST'])
    @login_required
    def importar_flexivel():
        uploaded = request.files.get('arquivo')
        if not uploaded or not uploaded.filename.lower().endswith('.xlsx'):
            return jsonify({'erro': 'Envie um arquivo .xlsx'}), 400
        if request.content_length and request.content_length > 10 * 1024 * 1024:
            return jsonify({'erro': 'Arquivo excede 10 MB'}), 413
        from openpyxl import load_workbook
        try:
            wb = load_workbook(uploaded, read_only=True, data_only=True)
        except Exception:
            return jsonify({'erro': 'Arquivo Excel invalido ou corrompido'}), 400
        selected = set(request.form.getlist('abas'))
        records, errors = [], []
        try:
            for ws in wb.worksheets:
                if selected and ws.title not in selected:
                    continue
                found, rejected = parse_sheet(ws, ws.title)
                records.extend(found)
                errors.extend(rejected)
        finally:
            wb.close()
        if len(records) > 20000:
            return jsonify({'erro': 'Limite total de 20000 registros excedido'}), 413
        db = get_db()
        existing = {key_of(dict(row)) for row in db.execute(
            'SELECT nome, especialidade, mes_ano FROM estagios').fetchall()}
        new, duplicates = [], 0
        for item in records:
            key = key_of(item)
            if key in existing:
                duplicates += 1
            else:
                new.append(item)
                existing.add(key)
        result = {'total_planilha': len(records), 'novos': len(new),
                  'duplicados': duplicates, 'erros': errors[:100],
                  'total_erros': len(errors), 'preview': new[:30],
                  'abas_disponiveis': wb.sheetnames}
        if not request.args.get('confirmar') == '1':
            return jsonify(result)
        if errors:
            return jsonify({**result, 'erro': 'Corrija as linhas com erro antes de confirmar'}), 422
        if not records:
            return jsonify({**result, 'erro': 'Nenhum registro valido encontrado'}), 422
        try:
            for rec in new:
                cur = db.execute(
                    'INSERT INTO estagios (tipo_id, mes_ano, semana, nome, especialidade, cracha, termino, email, telefone, observacao, etapa) '
                    'VALUES (:tipo_id, :mes_ano, :semana, :nome, :especialidade, :cracha, :termino, :email, :telefone, :observacao, :etapa)', rec)
                db.execute('INSERT INTO historico_etapas (estagio_id, etapa, observacao, responsavel) VALUES (?, ?, ?, ?)',
                           (cur.lastrowid, rec['etapa'], 'Importado via planilha flexivel', current_user.nome))
            db.commit()
        except Exception:
            db.rollback()
            return jsonify({'erro': 'Falha na gravacao. Nenhuma alteracao desta importacao foi confirmada.'}), 500
        return jsonify({**result, 'importados': len(new)})
